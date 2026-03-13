"""Debug: inspect actual cell structure of XLSX files."""
import sys
sys.path.insert(0, "backend")
import openpyxl

# ── Posição - Acoes sheet ──
print("=== POSIÇÃO - Acoes (detailed) ===")
wb = openpyxl.load_workbook("data/exemplob3/posicao-2026-03-12-14-23-06.xlsx", data_only=True)
ws = wb["Acoes"]
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

# Print first 5 rows with ALL columns
for row_idx in range(1, min(6, ws.max_row + 1)):
    cells = []
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=row_idx, column=col_idx).value
        cells.append(val)
    print(f"Row {row_idx}: {cells}")

print(f"\n--- Acoes: column headers at row 1 ---")
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print(headers)

print(f"\n--- Acoes: row 2 ---")
row2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
print(row2)
wb.close()

# ── Negociação ──
print("\n=== NEGOCIAÇÃO (detailed) ===")
wb = openpyxl.load_workbook("data/exemplob3/negociacao-2026-03-12-14-24-20.xlsx", data_only=True)
ws = wb["Negociação"]
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
for row_idx in range(1, min(4, ws.max_row + 1)):
    cells = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Row {row_idx}: {cells}")
wb.close()

# ── Movimentação ──
print("\n=== MOVIMENTAÇÃO (detailed) ===")
wb = openpyxl.load_workbook("data/exemplob3/movimentacao-2026-03-12-14-23-56.xlsx", data_only=True)
ws = wb["Movimentação"]
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
for row_idx in range(1, min(4, ws.max_row + 1)):
    cells = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"Row {row_idx}: {cells}")
wb.close()
