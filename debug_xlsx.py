"""Debug: inspect actual XLSX structure."""
import sys
sys.path.insert(0, "backend")
import openpyxl

# ── Posição ──
print("=== POSIÇÃO SHEETS ===")
wb = openpyxl.load_workbook("data/exemplob3/posicao-2026-03-12-14-23-06.xlsx", read_only=True, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = list(ws.iter_rows(max_row=3, values_only=True))
    print(f"\nSheet: '{sn}' (repr: {repr(sn)})")
    for i, r in enumerate(rows):
        print(f"  Row {i}: {r}")
wb.close()

# ── Negociação ──
print("\n=== NEGOCIAÇÃO SHEETS ===")
wb = openpyxl.load_workbook("data/exemplob3/negociacao-2026-03-12-14-24-20.xlsx", read_only=True, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = list(ws.iter_rows(max_row=3, values_only=True))
    print(f"\nSheet: '{sn}' (repr: {repr(sn)})")
    for i, r in enumerate(rows):
        print(f"  Row {i}: {r}")
wb.close()

# ── Movimentação ──
print("\n=== MOVIMENTAÇÃO SHEETS ===")
wb = openpyxl.load_workbook("data/exemplob3/movimentacao-2026-03-12-14-23-56.xlsx", read_only=True, data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = list(ws.iter_rows(max_row=3, values_only=True))
    print(f"\nSheet: '{sn}' (repr: {repr(sn)})")
    for i, r in enumerate(rows):
        print(f"  Row {i}: {r}")
wb.close()
